"""
Excel → CSV 変換ツール

Excelファイルの各シート（＝各現場）を読み取り、
financial_analyzer.py が読めるCSV形式に変換します。

【使い方】
  python excel_to_csv.py <Excelファイルのパス>

  例: python excel_to_csv.py data/全現場PL.xlsx

【Excelの条件】
  - 各シートが1現場のPLデータ（シート名＝現場名）
  - 列: 月ごとのデータ（2024年4月, 2025年4月, 2024年5月, ... の順）
  - A列に行ラベル（売上、出庫、タッチ数、...）がある
  - 行1に年、行2に月が記載されている
"""

import sys
import os
import csv
from openpyxl import load_workbook


# ============================================================
# A列のラベルと出力フィールドの対応
# Excelに書かれているラベル名 → CSVのフィールド名
# 部分一致で検索するので、多少の表記揺れに対応
# ============================================================

LABEL_MAP = [
    ("売上",                "revenue"),
    ("出庫",                "shipments"),
    ("タッチ数",            "touch_count"),
    ("平均単価",            "avg_unit_price"),
    ("管理費",              "admin_cost"),
    ("その他売上",          "other_revenue"),
    ("給与手当",            "salary_cost"),
    ("社員総工数",          "employee_total_hours"),
    ("社員通常工数",        "employee_regular_hours"),
    ("社員残業工数",        "employee_overtime_hours"),
    ("社員人数",            "employee_count"),
    ("労務費",              "labor_cost"),
    ("スタッフ工数",        "staff_hours"),
    ("平均時給",            "avg_hourly_wage"),
    ("外注費",              "outsource_cost"),
    ("タイミー工数",        "timee_hours"),
    ("タイミー時給",        "timee_hourly_cost"),
    ("有給",                "paid_leave"),
    ("法定福利費",          "statutory_welfare"),
    ("旅費交通費",          "travel_expense"),
]

FLOAT_FIELDS = {
    "avg_unit_price", "employee_total_hours", "employee_regular_hours",
    "employee_overtime_hours", "employee_count", "staff_hours", "timee_hours",
}


def safe_int(val) -> int:
    if val is None:
        return 0
    try:
        return int(float(str(val).replace(",", "").replace("¥", "").replace("\\", "").replace("円", "")))
    except (ValueError, TypeError):
        return 0


def safe_float(val) -> float:
    if val is None:
        return 0.0
    try:
        return float(str(val).replace(",", "").replace("¥", "").replace("\\", "").replace("円", ""))
    except (ValueError, TypeError):
        return 0.0


def detect_columns(ws) -> list[dict]:
    """
    ヘッダーから年・月の列ペアを検出する。
    行1〜5の範囲で年と月の情報を探す。
    """
    columns = []

    # 年の情報が入っている行と月の情報が入っている行を探す
    year_row = None
    month_row = None

    for row in range(1, 6):
        for col in range(2, min(ws.max_column + 1, 50)):
            val = ws.cell(row=row, column=col).value
            if val is None:
                continue
            s = str(val).strip()
            # 年の検出
            if year_row is None:
                for y in range(2020, 2030):
                    if str(y) in s:
                        year_row = row
                        break
            # 月の検出
            if month_row is None:
                for m in range(1, 13):
                    if f"{m}月" == s or s.endswith("月"):
                        month_row = row
                        break
        if year_row and month_row:
            break

    if not year_row or not month_row:
        return columns

    # 各列の年・月を取得
    current_year = None
    for col in range(2, min(ws.max_column + 1, 50)):
        # 年
        year_val = ws.cell(row=year_row, column=col).value
        if year_val is not None:
            s = str(year_val).strip()
            for y in range(2020, 2030):
                if str(y) in s:
                    current_year = y
                    break

        # 月
        month_val = ws.cell(row=month_row, column=col).value
        if month_val is not None and current_year is not None:
            s = str(month_val).strip()
            for m in range(1, 13):
                if f"{m}月" == s:
                    columns.append({"col": col, "year": current_year, "month": m})
                    break

    return columns


def detect_row_map(ws) -> dict[str, int]:
    """A列のラベルを読み取り、各フィールドの行番号を特定する"""
    row_map = {}
    used_labels = set()

    for row in range(1, min(ws.max_row + 1, 100)):
        cell_val = ws.cell(row=row, column=1).value
        if cell_val is None:
            continue

        label = str(cell_val).strip()

        for search_text, field_name in LABEL_MAP:
            if field_name in row_map:
                continue
            if search_text in label and search_text not in used_labels:
                row_map[field_name] = row
                used_labels.add(search_text)
                break

    return row_map


def read_sheet(ws, site_name: str, columns: list[dict], row_map: dict[str, int]) -> list[dict]:
    """1つのシートからデータを読み取る"""
    records = []

    for col_info in columns:
        col = col_info["col"]
        year = col_info["year"]
        month = col_info["month"]

        record = {
            "site_name": site_name,
            "year": year,
            "month": month,
        }

        for field_name in [f for _, f in LABEL_MAP]:
            if field_name not in row_map:
                # フィールドが見つからない場合はデフォルト値
                record[field_name] = 0.0 if field_name in FLOAT_FIELDS else 0
                continue

            val = ws.cell(row=row_map[field_name], column=col).value

            if field_name in FLOAT_FIELDS:
                record[field_name] = safe_float(val)
            else:
                record[field_name] = safe_int(val)

        if record["revenue"] > 0:
            records.append(record)

    return records


def convert_excel_to_csv(excel_path: str, output_path: str) -> int:
    """ExcelファイルをCSVに変換"""
    print(f"\n  Excelファイル読み込み: {excel_path}")

    wb = load_workbook(excel_path, data_only=True)
    print(f"  シート数: {len(wb.sheetnames)}")
    print(f"  シート名: {', '.join(wb.sheetnames)}")

    all_records = []
    fieldnames = ["site_name", "year", "month"] + [f for _, f in LABEL_MAP]

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]

        # ヘッダーから列情報を検出
        columns = detect_columns(ws)
        if not columns:
            print(f"  警告: '{sheet_name}' から年月情報を検出できませんでした。スキップ。")
            continue

        # A列のラベルから行番号を検出
        row_map = detect_row_map(ws)
        if not row_map:
            print(f"  警告: '{sheet_name}' から行ラベルを検出できませんでした。スキップ。")
            continue

        # 検出結果の表示（最初のシートのみ）
        if not all_records:
            print(f"\n  検出された行マッピング（{sheet_name}）:")
            for search_text, field_name in LABEL_MAP:
                if field_name in row_map:
                    print(f"    行{row_map[field_name]:>3d}: {search_text} → {field_name}")
                else:
                    print(f"    ???: {search_text} → 未検出")
            print()

        records = read_sheet(ws, sheet_name, columns, row_map)
        all_records.extend(records)
        print(f"  {sheet_name}: {len(columns)}列検出, {len(records)}件のデータを読み取り")

    if not all_records:
        print("\n  エラー: データが1件も読み取れませんでした。")
        print("  以下を確認してください:")
        print("    1. 各シートのA列に「売上」「出庫」等のラベルがあるか")
        print("    2. 行1-2に「2024年」「4月」等のヘッダーがあるか")
        print("    3. シート名が現場名になっているか")
        return 0

    with open(output_path, "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(all_records)

    sites = set(r["site_name"] for r in all_records)
    print(f"\n  CSV出力完了: {output_path}")
    print(f"  合計: {len(all_records)}件 / {len(sites)}現場")

    return len(all_records)


def main():
    if len(sys.argv) < 2:
        print()
        print("  ╔═══════════════════════════════════════════════╗")
        print("  ║   Excel → CSV 変換ツール                      ║")
        print("  ╚═══════════════════════════════════════════════╝")
        print()
        print("  使い方:")
        print("    python excel_to_csv.py <Excelファイルのパス>")
        print()
        print("  例:")
        print("    python excel_to_csv.py data/全現場PL.xlsx")
        print("    python excel_to_csv.py /home/user/Documents/現場データ.xlsx")
        print()
        print("  変換後:")
        print("    python financial_analyzer.py")
        print("  で全現場の分析が実行されます。")
        return

    excel_path = sys.argv[1]

    if not os.path.exists(excel_path):
        print(f"\n  エラー: ファイルが見つかりません: {excel_path}")
        return

    base_dir = os.path.dirname(os.path.abspath(__file__))
    output_path = os.path.join(base_dir, "data", "site_financial_data.csv")

    count = convert_excel_to_csv(excel_path, output_path)

    if count > 0:
        print(f"\n  次のステップ:")
        print(f"    python financial_analyzer.py")
        print(f"  を実行すると、全{count}件のデータで分析が行われます。")


if __name__ == "__main__":
    main()
